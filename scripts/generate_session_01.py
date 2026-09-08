#!/usr/bin/env python3
"""
HashCode Presentation Exporter
Generates Session #01 as:
  - PowerPoint (.pptx)
  - PDF (.pdf)

Dependencies:
  pip install python-pptx reportlab
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from reportlab.lib.pagesizes import landscape
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "dist"
OUTPUT.mkdir(exist_ok=True)

PPTX_OUT = OUTPUT / "hashcode-session-01-ia-productivite-ai-agents.pptx"
PDF_OUT = OUTPUT / "hashcode-session-01-ia-productivite-ai-agents.pdf"

W, H = 13.333, 7.5
LIME = RGBColor(197, 244, 65)
BG = RGBColor(10, 12, 15)
WHITE = RGBColor(242, 244, 245)
MUTED = RGBColor(145, 153, 162)
DARK = RGBColor(22, 26, 31)

SLIDES = [
    {
        "type": "cover",
        "eyebrow": "HASHCODE WORKSHOP #01",
        "title": "IA, Productivité\net AI Agents",
        "subtitle": "Comprendre comment l'intelligence artificielle peut améliorer les processus de travail et soutenir la prise d'action.",
        "meta": "Workshop interactif  |  Comprendre · Analyser · Expérimenter · Produire"
    },
    {
        "type": "statement",
        "eyebrow": "01 / QUESTION CENTRALE",
        "title": "Sommes-nous réellement\nplus productifs ?",
        "subtitle": "Ou sommes-nous principalement davantage sollicités ?"
    },
    {
        "type": "cards",
        "eyebrow": "02 / LE PROBLÈME",
        "title": "Une part importante du travail\nest désormais consacrée à la friction.",
        "cards": [
            ("01", "Recherche dispersée", "L'information existe, mais elle est répartie entre plusieurs outils et conversations."),
            ("02", "Tâches répétitives", "Les mêmes opérations de lecture, synthèse, rédaction et organisation reviennent constamment."),
            ("03", "Fragmentation de l'attention", "Les changements fréquents de contexte réduisent la concentration et ralentissent la production.")
        ],
        "footer": "Être occupé n'est pas la même chose qu'être productif."
    },
    {
        "type": "cards",
        "eyebrow": "03 / RÉSULTATS ATTENDUS",
        "title": "Résultats attendus\nà l'issue de la séance.",
        "cards": [
            ("01", "Identifier", "Repérer les principales frictions dans votre propre travail."),
            ("02", "Comprendre", "Distinguer les usages réellement utiles de l'IA des simples effets de mode."),
            ("03", "Concevoir", "Imaginer un workflow augmenté et une première idée d'AI Agent.")
        ],
        "footer": "Comprendre → Analyser → Expérimenter → Produire"
    },
    {
        "type": "split",
        "eyebrow": "04 / MODÈLE MENTAL",
        "title": "Humain × IA",
        "left": ("HUMAIN", ["Jugement", "Contexte", "Responsabilité"]),
        "right": ("IA", ["Vitesse", "Analyse", "Première version"]),
        "result": "CAPACITÉ\nAUGMENTÉE",
        "footer": "L'IA ne remplace pas la pensée. Elle amplifie un processus déjà compris."
    },
    {
        "type": "flow",
        "eyebrow": "05 / FRAMEWORK HASHCODE",
        "title": "Avant d'automatiser,\nil faut comprendre le travail.",
        "steps": [
            ("01", "Identifier", "Où perdez-vous réellement du temps ou de l'énergie ?"),
            ("02", "Simplifier", "Le processus actuel est-il vraiment nécessaire ?"),
            ("03", "Augmenter", "Comment l'IA peut-elle améliorer la réflexion ou la production ?"),
            ("04", "Automatiser", "Quelles étapes répétitives peuvent être déléguées dans un cadre contrôlé ?")
        ],
        "footer": "Automatiser un mauvais processus permet simplement de produire des erreurs plus vite."
    },
    {
        "type": "comparison",
        "eyebrow": "06 / DISTINCTION ESSENTIELLE",
        "title": "Assistant IA et AI Agent",
        "left": ("ASSISTANT IA", "Interaction assistée", ["Interaction directe", "Réponse à la demande", "Pilotage humain explicite"]),
        "right": ("AI AGENT", "Exécution orientée objectif", ["Objectif défini", "Processus multi-étapes", "Outils et actions autorisés"]),
        "footer": "La distinction porte sur le niveau de délégation, l'utilisation d'outils et le degré d'autonomie autorisé."
    },
    {
        "type": "flow",
        "eyebrow": "07 / ARCHITECTURE D'UN AGENT",
        "title": "Les composants essentiels\nd'un AI Agent.",
        "steps": [
            ("01", "Objectif", "Définir le résultat attendu."),
            ("02", "Raisonnement", "Choisir la prochaine étape."),
            ("03", "Outils", "Accéder aux capacités autorisées."),
            ("04", "Action", "Produire un effet réel.")
        ],
        "footer": "Observer → Ajuster → Continuer ou s'arrêter. Les permissions et validations sont essentielles."
    },
    {
        "type": "flow",
        "eyebrow": "08 / DÉMONSTRATION",
        "title": "Transformer un workflow réel.",
        "steps": [
            ("01", "Observer", "Décrire le processus actuel."),
            ("02", "Identifier", "Repérer la friction principale."),
            ("03", "Augmenter", "Utiliser l'IA sur une étape précise."),
            ("04", "Mesurer", "Comparer le processus avant et après.")
        ],
        "footer": "L'objectif est d'améliorer un problème réel, pas d'ajouter de l'IA partout."
    },
    {
        "type": "cards",
        "eyebrow": "09 / WORKSHOP",
        "title": "AI Productivity Case",
        "cards": [
            ("A", "Problème", "Quelle tâche ou difficulté souhaitez-vous améliorer ?"),
            ("B", "Processus", "Comment cette activité fonctionne-t-elle aujourd'hui ?"),
            ("C", "Amélioration", "Quelle étape l'IA peut-elle augmenter sans supprimer le contrôle nécessaire ?")
        ],
        "footer": "Livrable : un workflow Before → After."
    },
    {
        "type": "cards",
        "eyebrow": "10 / WORKSHOP",
        "title": "AI Agent Idea Card",
        "cards": [
            ("01", "Mission", "Quel problème précis l'agent doit-il résoudre ?"),
            ("02", "Capacités", "Quelles informations, outils et actions sont nécessaires ?"),
            ("03", "Contrôle", "Quelles limites et validations humaines doivent être imposées ?")
        ],
        "footer": "Livrable : une première spécification d'AI Agent."
    },
    {
        "type": "statement",
        "eyebrow": "11 / SYNTHÈSE",
        "title": "Identifier la friction.\nSimplifier le processus.\nAugmenter avec l'IA.",
        "subtitle": "Automatiser uniquement ce qui est suffisamment compris et contrôlé."
    },
    {
        "type": "statement",
        "eyebrow": "12 / NEXT ACTION",
        "title": "Choisissez un seul workflow.",
        "subtitle": "Analysez-le. Simplifiez-le. Puis améliorez-le avec l'IA.\n\nSmall workflow. Real problem. Measurable impact."
    }
]


def add_text(slide, x, y, w, h, text, size, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def background(slide, prs):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(0.08)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = LIME
    slide.shapes[-1].line.fill.background()
    add_text(slide, 11.8, 7.05, 1.0, 0.2, "HASHCODE", 7, MUTED, True, PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, number, title, body):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK
    shape.line.color.rgb = RGBColor(55, 61, 68)
    add_text(slide, x+.28, y+.25, w-.5, .3, number, 9, LIME, True)
    add_text(slide, x+.28, y+.78, w-.5, .55, title, 17, WHITE, True)
    add_text(slide, x+.28, y+1.45, w-.55, h-1.7, body, 10, MUTED)


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    for data in SLIDES:
        slide = prs.slides.add_slide(blank)
        background(slide, prs)
        add_text(slide, .7, .55, 8, .3, data.get("eyebrow", ""), 9, LIME, True)

        if data["type"] == "cover":
            add_text(slide, .8, 1.35, 9.5, 1.7, data["title"], 36, WHITE, True)
            add_text(slide, .8, 3.35, 7.6, 1.0, data["subtitle"], 17, MUTED)
            add_text(slide, .8, 5.75, 9, .45, data["meta"], 10, LIME, True)
            # Abstract signal
            shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.0), Inches(2.0), Inches(2.1), Inches(2.1))
            shape.fill.background()
            shape.line.color.rgb = LIME
            add_text(slide, 10.0, 2.7, 2.1, .5, "AI", 24, LIME, True, PP_ALIGN.CENTER)

        elif data["type"] == "statement":
            add_text(slide, .9, 1.65, 10.8, 2.2, data["title"], 31, WHITE, True)
            add_text(slide, .95, 4.5, 9.8, 1.1, data["subtitle"], 17, MUTED)

        elif data["type"] == "cards":
            add_text(slide, .8, 1.1, 11.5, 1.1, data["title"], 28, WHITE, True)
            cards = data["cards"]
            cw = 3.75
            for i, card in enumerate(cards):
                add_card(slide, .8+i*4.1, 2.55, cw, 2.7, *card)
            add_text(slide, .9, 5.85, 11.2, .5, data["footer"], 13, WHITE, True)

        elif data["type"] == "split":
            add_text(slide, .8, 1.15, 7.5, .8, data["title"], 28, WHITE, True)
            for x, block in [(1.0, data["left"]), (5.0, data["right"])]:
                label, items = block
                add_card(slide, x, 2.2, 3.1, 2.8, label, label, "\n".join(items))
            add_card(slide, 9.2, 2.2, 3.0, 2.8, "RESULT", data["result"], "Humain et IA combinent leurs capacités.")
            add_text(slide, 1.0, 5.75, 11, .5, data["footer"], 13, WHITE, True)

        elif data["type"] == "flow":
            add_text(slide, .8, 1.05, 11.7, 1.05, data["title"], 27, WHITE, True)
            steps = data["steps"]
            for i, step in enumerate(steps):
                x = .55 + i*3.22
                add_card(slide, x, 2.55, 2.75, 2.45, *step)
                if i < len(steps)-1:
                    add_text(slide, x+2.78, 3.55, .3, .4, "→", 18, LIME, True, PP_ALIGN.CENTER)
            add_text(slide, .8, 5.75, 11.4, .55, data["footer"], 12, WHITE, True)

        elif data["type"] == "comparison":
            add_text(slide, .8, 1.1, 11.5, .8, data["title"], 28, WHITE, True)
            left = data["left"]
            right = data["right"]
            add_card(slide, .9, 2.25, 5.25, 3.2, left[0], left[1], "\n".join("• "+x for x in left[2]))
            add_card(slide, 7.15, 2.25, 5.25, 3.2, right[0], right[1], "\n".join("• "+x for x in right[2]))
            add_text(slide, 6.25, 3.55, .8, .5, "→", 24, LIME, True, PP_ALIGN.CENTER)
            add_text(slide, .9, 5.9, 11.2, .55, data["footer"], 11, MUTED)

    prs.save(PPTX_OUT)


def pdf_text(c, text, x, y, size=18, color=(1,1,1), leading=None):
    c.setFillColorRGB(*color)
    c.setFont("Helvetica-Bold" if size >= 18 else "Helvetica", size)
    leading = leading or size*1.25
    for line in text.split("\n"):
        c.drawString(x, y, line)
        y -= leading
    return y


def build_pdf():
    pw, ph = landscape((960, 540))
    c = canvas.Canvas(str(PDF_OUT), pagesize=(pw, ph))

    for data in SLIDES:
        c.setFillColorRGB(10/255, 12/255, 15/255)
        c.rect(0, 0, pw, ph, fill=1, stroke=0)
        c.setFillColorRGB(197/255, 244/255, 65/255)
        c.rect(0, ph-6, pw, 6, fill=1, stroke=0)

        y = ph-55
        y = pdf_text(c, data.get("eyebrow",""), 45, y, 9, (197/255,244/255,65/255))
        y -= 15
        title = data.get("title","")
        y = pdf_text(c, title, 45, y, 28, (242/255,244/255,245/255))

        if data["type"] == "statement":
            pdf_text(c, data["subtitle"], 45, 170, 16, (.6,.63,.67))
        elif data["type"] == "cover":
            pdf_text(c, data["subtitle"], 45, 250, 16, (.6,.63,.67))
            pdf_text(c, data["meta"], 45, 95, 10, (197/255,244/255,65/255))
        elif data["type"] in ("cards","flow"):
            items = data.get("cards") or data.get("steps")
            x = 45
            for item in items:
                c.setFillColorRGB(22/255,26/255,31/255)
                c.roundRect(x, 150, 195, 180, 8, fill=1, stroke=0)
                pdf_text(c, item[0], x+15, 305, 9, (197/255,244/255,65/255))
                pdf_text(c, item[1], x+15, 270, 15, (1,1,1))
                pdf_text(c, item[2], x+15, 230, 9, (.6,.63,.67))
                x += 220
            pdf_text(c, data["footer"], 45, 90, 11, (1,1,1))
        else:
            pdf_text(c, data.get("footer",""), 45, 100, 12, (1,1,1))

        c.showPage()

    c.save()


if __name__ == "__main__":
    build_pptx()
    build_pdf()
    print(f"Generated: {PPTX_OUT}")
    print(f"Generated: {PDF_OUT}")
